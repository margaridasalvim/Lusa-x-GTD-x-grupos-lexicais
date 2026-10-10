# -*- coding: utf-8 -*-
"""
LUSA ANALISE — LUSA × GTD — streamlit_app.py

Inclui:
- Só Lusa (fixo): lower(fonte) LIKE '%lusa%'
- País principal na UI: place_country (legível). ISO (place_country_iso) só interno para GTD
- Filtro 1: Action type (principal) = action_type
- 2.º filtro: Violência (detalhe) por título
- Autor macro (por título)
- Política (Nacional / Externa / qualquer) + modo (Todos / Só / Excluir)
- Famílias lexicais com normalização robusta
- Pesquisa de palavras com modo Texto (contém) e Regex
  (pesquisa sempre em título + keyword_text_sources)
- Tabela Lusa (Top N)
- Export Lusa:
  * CSV completo
  * Excel legível com ano + título + place_country
- Sidebar:
  * Países agrupados no fundo
  * Top 20 países atuais com contagens
  * seletor com países ordenados por frequência
- Tabela de países mais referidos — Lusa
- Export de países mais referidos:
  * CSV
  * Excel com dados + gráfico
- Secção experimental:
  * deteção automática de países no título
  * tabela separada
  * export CSV / Excel
- Comparação Lusa × GTD por dia+ISO
- GTD event-level “lista completa” (gtd_raw) — preview + export CSV
- Gráfico temporal: Lusa + GTD
- Export gráfico para Excel:
  * folha "dados" em formato largo: ano | Lusa | GTD
  * folha com gráfico Excel já criado

Alterações nesta versão:
- heurística de violência/ataque mais robusta
- "Consumado" deixa de depender apenas de mortos/feridos/explosão
- títulos com sinais claros de ataque (atentado, ataque, carro-bomba, abriu fogo, etc.)
  passam a contar como "Consumado" quando não houver marcadores de falha/tentativa/plano
- coluna auxiliar attack_detected para auditoria interna/visual
"""

import re
import unicodedata
from io import BytesIO
from pathlib import Path

import altair as alt
import duckdb
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

# --------------------------------------------------
# CONFIG
# --------------------------------------------------
from PIL import Image

st.set_page_config(
    page_title="Lusa análise",
    page_icon=Image.open("logo_lxt.jpg"),
    layout="wide"
)

# INICIO EXCERTO RELATORIO DE MESTRADO

# Estilo visual do botão de informações
st.markdown("""
<style>
div.st-key-btn_info_relatorio button {
    background-color: #B22222 !important;
    color: white !important;
    border: 1px solid #B22222 !important;
    font-weight: 600 !important;
}
div.st-key-btn_info_relatorio button:hover {
    background-color: #8B1A1A !important;
    color: white !important;
    border-color: #8B1A1A !important;
}
</style>
""", unsafe_allow_html=True)

if "mostrar_info_relatorio" not in st.session_state:
    st.session_state["mostrar_info_relatorio"] = False

_rotulo_info = (
    "Informações sobre o projeto"
    if st.session_state["mostrar_info_relatorio"]
    else "Informações sobre o projeto"
)

if st.button(_rotulo_info, key="btn_info_relatorio"):
    st.session_state["mostrar_info_relatorio"] = (
        not st.session_state["mostrar_info_relatorio"]
    )

if st.session_state["mostrar_info_relatorio"]:
    st.markdown(
        "**Projeto desenvolvido em 2026, no âmbito do relatorio "
        "de mestrado em Jornalismo da NOVA FCSH.**"
    )
    st.markdown(
        "O texto seguinte foi retirado do relatório de mestrado "
        "no âmbito do qual foi desenvolvido este projeto."
    )

    _ficheiro_relatorio = Path(__file__).resolve().parent / "Texto_relatorio_mestrado.md"

    if _ficheiro_relatorio.is_file():
        _conteudo_relatorio = _ficheiro_relatorio.read_text(encoding="utf-8")

        # Transformar os subtítulos numerados em títulos encarnados,
        # apenas na apresentação da aplicação.
        _conteudo_relatorio = re.sub(
            r"(?ms)^\*\*2\.\d+\.?\s*(.*?)\*\*",
            lambda m: '<h3 style="color: #B22222; font-weight: 700;">'
            + " ".join(m.group(1).split())
            + "</h3>",
            _conteudo_relatorio,
        )

        st.markdown(_conteudo_relatorio, unsafe_allow_html=True)
    else:
        st.warning("Não foi encontrado o ficheiro Texto_relatorio_mestrado.md.")

# FIM EXCERTO RELATORIO DE MESTRADO

PROJECT_ROOT = Path(__file__).resolve().parent
DB_PATH = PROJECT_ROOT / "db" / "analysis.duckdb"

USE_CONSERVATIVE_VIOLENCE_FALLBACK = True

# Aliases experimentais para deteção de países no título
COUNTRY_TITLE_ALIASES = {
    "Estados Unidos": ["estados unidos", "eua", "e.u.a.", "washington"],
    "Reino Unido": [
        "reino unido", "britanico", "britânica", "britanica",
        "britanicos", "britânicos", "londres"
    ],
    "França": ["franca", "frances", "francês", "francesa", "franceses", "paris"],
    "Alemanha": ["alemanha", "alemao", "alemão", "alema", "alemã", "alemaes", "alemães", "berlim"],
    "Portugal": ["portugal", "portugues", "português", "portuguesa", "portugueses", "lisboa"],
    "Espanha": ["espanha", "espanhol", "espanhola", "madrid"],
    "Rússia": ["russia", "russo", "russa", "russos", "russas", "moscovo", "moscou"],
    "Ucrânia": ["ucrania", "ucraniano", "ucraniana", "ucranianos", "kiev", "kyiv"],
    "Síria": ["siria", "sirio", "sírio", "damasco"],
    "Israel": ["israel", "israelita", "israelitas", "tel avive", "telavive", "jerusalem", "jerusalém"],
    "Palestina": ["palestina", "palestiniano", "palestiniana", "gaza", "cisjordania", "cisjordânia"],
    "Afeganistão": ["afeganistao", "afeganistão", "afegao", "afegão", "cabul"],
    "Iraque": ["iraque", "iraquiano", "iraquiana", "bagdade"],
    "Irão": ["irao", "irão", "iraniano", "iraniana", "teera", "teerao", "teerão"],
    "Turquia": ["turquia", "turco", "turca", "ancara", "ankara", "istambul"],
}


@st.cache_resource
def get_con():
    return duckdb.connect(str(DB_PATH), read_only=True)


con = get_con()

# --------------------------------------------------
# Helpers gerais
# --------------------------------------------------
def parse_date_series(s: pd.Series) -> pd.Series:
    t1 = pd.to_datetime(s, errors="coerce")
    if t1.notna().mean() > 0.95:
        return t1
    t2 = pd.to_datetime(s, errors="coerce", dayfirst=True)
    return t2 if t2.notna().mean() >= t1.notna().mean() else t1


def safe_int_series(s: pd.Series) -> pd.Series:
    return pd.to_numeric(s, errors="coerce").fillna(0).astype(int)


def normalize_text(x) -> str:
    if pd.isna(x):
        return ""
    x = str(x).strip().lower()
    x = unicodedata.normalize("NFKD", x)
    x = "".join(c for c in x if not unicodedata.combining(c))
    return x


def normalize_series(s: pd.Series) -> pd.Series:
    s = s.fillna("").astype(str).str.strip().str.lower()
    return s.map(
        lambda x: "".join(
            c for c in unicodedata.normalize("NFKD", x)
            if not unicodedata.combining(c)
        )
    )


def ensure_columns(df: pd.DataFrame, cols_with_defaults: dict) -> pd.DataFrame:
    for col, default in cols_with_defaults.items():
        if col not in df.columns:
            df[col] = default
    return df


def df_to_excel_bytes(df: pd.DataFrame, sheet_name: str = "dados") -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
    return output.getvalue()


def df_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")


def auto_fit_columns(ws, max_width: int = 45):
    for col in ws.columns:
        length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                length = max(length, len(str(cell.value)) if cell.value is not None else 0)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), max_width)


def build_chart_excel_bytes(df_wide: pd.DataFrame, title: str = "Evolução temporal — Lusa vs GTD") -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_wide.to_excel(writer, index=False, sheet_name="dados")

    output.seek(0)
    wb = load_workbook(output)
    ws = wb["dados"]
    auto_fit_columns(ws)

    ws_chart = wb.create_sheet("grafico")
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = "Ocorrências"
    chart.x_axis.title = "Ano"
    chart.height = 10
    chart.width = 18
    chart.style = 2

    max_row = ws.max_row
    max_col = ws.max_column

    if max_row >= 2 and max_col >= 2:
        data = Reference(ws, min_col=2, max_col=max_col, min_row=1, max_row=max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws_chart.add_chart(chart, "B2")

    final_output = BytesIO()
    wb.save(final_output)
    return final_output.getvalue()


def build_country_summary(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty or "place_country" not in df.columns:
        return pd.DataFrame(columns=["place_country", "noticias", "percentagem"])

    out = (
        df.copy()
        .assign(place_country=lambda x: x["place_country"].fillna("").astype(str).str.strip())
        .loc[lambda x: x["place_country"] != ""]
        .groupby("place_country", as_index=False)
        .size()
        .rename(columns={"size": "noticias"})
        .sort_values(["noticias", "place_country"], ascending=[False, True])
    )

    if out.empty:
        out["percentagem"] = []
        return out

    total = int(out["noticias"].sum())
    out["percentagem"] = (out["noticias"] / total * 100).round(2)
    return out


def build_country_excel_bytes(df_country: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_country.to_excel(writer, index=False, sheet_name="paises")

    output.seek(0)
    wb = load_workbook(output)
    ws = wb["paises"]
    auto_fit_columns(ws)

    if ws.max_row >= 2:
        chart = LineChart()
        chart.title = "Países mais referidos"
        chart.y_axis.title = "Número de notícias"
        chart.x_axis.title = "País"
        chart.height = 10
        chart.width = 18
        chart.style = 2

        data = Reference(ws, min_col=2, max_col=2, min_row=1, max_row=ws.max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        ws_chart = wb.create_sheet("grafico_paises")
        ws_chart.add_chart(chart, "B2")

    final_output = BytesIO()
    wb.save(final_output)
    return final_output.getvalue()


def build_detected_country_excel_bytes(df_detected: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_detected.to_excel(writer, index=False, sheet_name="deteccao_titulo")

    output.seek(0)
    wb = load_workbook(output)
    ws = wb["deteccao_titulo"]
    auto_fit_columns(ws)

    if ws.max_row >= 2:
        chart = LineChart()
        chart.title = "Países detetados automaticamente no título"
        chart.y_axis.title = "Número de títulos"
        chart.x_axis.title = "País detetado"
        chart.height = 10
        chart.width = 18
        chart.style = 2

        data = Reference(ws, min_col=2, max_col=2, min_row=1, max_row=ws.max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        ws_chart = wb.create_sheet("grafico_deteccao")
        ws_chart.add_chart(chart, "B2")

    final_output = BytesIO()
    wb.save(final_output)
    return final_output.getvalue()

# --------------------------------------------------
# Autor / Tipologia macro (por título)
# --------------------------------------------------
AUTHOR_OPTIONS = [
    "Judicial",
    "Estado — Governo",
    "Estado — Forças Armadas",
    "Estado — Forças de Segurança",
    "Grupo não estatal",
    "ONG / Sociedade civil",
    "Universidades / Instituições académicas",
    "Indivíduo",
    "Autor omitido / evento",
]


def macro_author_from_title(title: str) -> str:
    t = title or ""

    if (
        "tribunal" in t or "supremo" in t or "juiz" in t or
        "ministerio publico" in t or "procurador" in t
    ):
        return "Judicial"

    if (
        "forcas armadas" in t or "exercito" in t or
        "marinha" in t or "forca aerea" in t
    ):
        return "Estado — Forças Armadas"

    if (
        "policia" in t or "psp" in t or "gnr" in t or
        "servicos de seguranca" in t or
        "servico secreto" in t or
        "inteligencia" in t
    ):
        return "Estado — Forças de Segurança"

    if (
        "governo" in t or "ministro" in t or "ministra" in t or
        "primeiro-ministro" in t or "primeiro ministro" in t or
        "parlamento" in t or "presidente" in t or "conselho de ministros" in t
    ):
        return "Estado — Governo"

    if (
        "universidade" in t or "politecnico" in t or
        "instituto" in t or "faculdade" in t or
        "centro de investigacao" in t or "academia" in t
    ):
        return "Universidades / Instituições académicas"

    if (
        (" ong " in f" {t} " or t.startswith("ong ") or t.endswith(" ong")) or
        ("organiza" in t and "nao governamental" in t) or
        "associacao" in t or "plataforma" in t or "movimento" in t or
        "coletivo" in t or "colectivo" in t or "amnesty" in t or "human rights watch" in t or "hrw" in t
    ):
        return "ONG / Sociedade civil"

    if (
        "milicia" in t or "grupo armado" in t or "grupos armados" in t or
        "organizacao terrorista" in t or
        "estado islamico" in t or "al-qaeda" in t or "talib" in t or "jihad" in t
    ):
        return "Grupo não estatal"

    if "suspeito" in t or "atacante" in t or "autor" in t:
        return "Indivíduo"

    return "Autor omitido / evento"

# --------------------------------------------------
# Política (por título)
# --------------------------------------------------
POLITICA_NAC_PAT = re.compile(
    r"\b(governo|parlamento|assembleia(?: da rep[úu]blica)?|presidente|primeiro-?ministro|"
    r"ministro|ministra|secret[aá]rio de estado|elei[cç][õo]es|campanha|"
    r"legislativas|presidenciais|aut[áa]rquicas|"
    r"ps\b|psd\b|cds\b|pcp\b|be\b|chega\b|il\b|livre\b|pan\b)\b",
    re.IGNORECASE
)

POLITICA_EXT_PAT = re.compile(
    r"\b(diplomacia|diplom[áa]tic\w*|pol[íi]tica externa|negocia[cç][ãa]o|cimeira|"
    r"uni[aã]o europeia|ue\b|comiss[aã]o europeia|conselho europeu|parlamento europeu|"
    r"onu\b|na[cç][õo]es unidas|nato\b|otan\b|g7\b|g20\b|brics\b|"
    r"san[cç][õo]es|embaixad\w*|consulado|minist[eé]rio dos neg[oó]cios estrangeiros|"
    r"mne\b|acordo|tratado|cessar-?fogo|processo de paz|mediac[aã]o)\b",
    re.IGNORECASE
)

# --------------------------------------------------
# Famílias lexicais
# --------------------------------------------------
FAM_PATTERNS = {
    "terror": re.compile(r"\bterror\w*", re.IGNORECASE),
    "radic": re.compile(r"\bradic\w*", re.IGNORECASE),
    "extrem": re.compile(r"\bextrem\w*", re.IGNORECASE),
    "rebel": re.compile(r"\brebel\w*", re.IGNORECASE),
}


def add_family_columns(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df

    df = ensure_columns(df, {"titulo": "", "keyword_text_sources": ""})

    if "titulo_norm" not in df.columns:
        df["titulo_norm"] = normalize_series(df["titulo"])

    if "keyword_text_sources_norm" not in df.columns:
        df["keyword_text_sources_norm"] = normalize_series(df["keyword_text_sources"])

    if "blob_title_kw_norm" not in df.columns:
        df["blob_title_kw_norm"] = df["titulo_norm"] + " " + df["keyword_text_sources_norm"]

    df["has_terror_title"] = df["titulo_norm"].str.contains(FAM_PATTERNS["terror"], na=False, regex=True).astype("int8")
    df["has_radic_title"] = df["titulo_norm"].str.contains(FAM_PATTERNS["radic"], na=False, regex=True).astype("int8")
    df["has_extrem_title"] = df["titulo_norm"].str.contains(FAM_PATTERNS["extrem"], na=False, regex=True).astype("int8")
    df["has_rebel_title"] = df["titulo_norm"].str.contains(FAM_PATTERNS["rebel"], na=False, regex=True).astype("int8")

    df["has_terror_kw"] = df["blob_title_kw_norm"].str.contains(FAM_PATTERNS["terror"], na=False, regex=True).astype("int8")
    df["has_radic_kw"] = df["blob_title_kw_norm"].str.contains(FAM_PATTERNS["radic"], na=False, regex=True).astype("int8")
    df["has_extrem_kw"] = df["blob_title_kw_norm"].str.contains(FAM_PATTERNS["extrem"], na=False, regex=True).astype("int8")
    df["has_rebel_kw"] = df["blob_title_kw_norm"].str.contains(FAM_PATTERNS["rebel"], na=False, regex=True).astype("int8")

    return df


def build_family_labels(df: pd.DataFrame, scope: str = "title") -> pd.Series:
    if scope == "title":
        terror = df["has_terror_title"] == 1
        radic = df["has_radic_title"] == 1
        extrem = df["has_extrem_title"] == 1
        rebel = df["has_rebel_title"] == 1
    else:
        terror = df["has_terror_kw"] == 1
        radic = df["has_radic_kw"] == 1
        extrem = df["has_extrem_kw"] == 1
        rebel = df["has_rebel_kw"] == 1

    out = pd.Series("—", index=df.index, dtype="object")

    for label, mask in [
        ("terror*", terror),
        ("radic*", radic),
        ("extrem*", extrem),
        ("rebel*", rebel),
    ]:
        out = out.mask(mask & (out == "—"), label)
        out = out.mask(
            mask & (out != "—") & ~out.str.contains(re.escape(label), regex=True),
            out + "; " + label
        )

    return out

# --------------------------------------------------
# Violência / Ataque (tipo + outcome) — por título
# --------------------------------------------------
VIOLENCE_TYPE_OPTIONS = [
    "Bomba/Explosão", "Tiroteio", "Ataque armado", "Esfaqueamento", "Sequestro/Rapto",
    "Atropelamento/Veículo", "Incêndio/Arson", "Granada/Morteiro", "Ataque químico", "Ataque suicida",
    "Outro/Não especificado"
]
VIOLENCE_OUTCOME_OPTIONS = ["Falhado/Tentativa", "Consumado", "Ameaça/Plano", "Desconhecido"]

VIOLENCE_TYPE_PATTERNS = [
    ("Bomba/Explosão", r"\b(bomba|bombista|explos[aã]o|explod\w*|detona\w*|deflagr\w*|artefacto explosivo|engenho explosivo|carro-bomba)\b"),
    ("Tiroteio", r"\b(tiroteio|tiros|disparos|abriu fogo|abrir fogo|rajad\w*|metralh\w*|balea\w*)\b"),
    ("Ataque armado", r"\b(ataque armado|homens armados|grupo armado|grupos armados|milician\w*|emboscad\w*|incurs[aã]o armada|assalto armado|atac\w*)\b"),
    ("Esfaqueamento", r"\b(esfaquead\w*|facad\w*|ataque (?:com )?faca)\b"),
    ("Sequestro/Rapto", r"\b(sequest\w*|rapt\w*|ref[eé]m|ref[eé]ns|tomar ref[eé]ns|tomada de ref[eé]ns)\b"),
    ("Atropelamento/Veículo", r"\b(atropel\w*|ataque com (?:carro|viatura|camioneta|cami[aã]o)|lan[çc]ou(?:-se)? com (?:carro|camioneta|cami[aã]o)|ve[ií]culo contra)\b"),
    ("Incêndio/Arson", r"\b(inc[eê]ndi\w*|fogo posto|atear fogo|arson)\b"),
    ("Granada/Morteiro", r"\b(granad\w*|morteir\w*|rocket|foguet\w*|m[ií]ssil|misseis|proj[eé]til)\b"),
    ("Ataque químico", r"\b(g[aá]s|sarin|cloro|qu[ií]mic\w*|ataque qu[ií]mic\w*)\b"),
    ("Ataque suicida", r"\b(suicid\w*|homem-bomba|mulher-bomba|bombista suicida)\b"),
]

# Marcadores de falha/tentativa
OUTCOME_FAILED_PAT = (
    r"\b("
    r"falh\w*|frustrad\w*|abortad\w*|impedid\w*|evitad\w*|desmantelad\w*|neutralizad\w*|"
    r"intercetad\w*|interceptad\w*|trav\w*|desativad\w*|desarmad\w*"
    r")\b"
)
OUTCOME_ATTEMPT_PAT = r"\b(tentativ\w*|tentou|tenta|tentaram|tentativa de)\b"

# Marcadores de ameaça/plano
OUTCOME_THREAT_PLAN_PAT = (
    r"\b("
    r"amea[cç]a\w*|amea[cç]ou|amea[çc]aram|plano\w*|planea\w*|planej\w*|conspira\w*|"
    r"prepara\w*|preparava\w*|pretendia\w*|pretens[aã]o de|alerta de ataque|"
    r"ameaça de ataque|ameaca de ataque|risco de ataque|suspeita de plano|"
    r"projeto de atentado"
    r")\b"
)

# Marcadores fortes de evento ocorrido
OUTCOME_OCCURRED_STRICT_PAT = (
    r"\b("
    r"explos[aã]o|explod\w*|deflagr\w*|rebent\w*|"
    r"matou|mataram|morto\w*|morte\w*|ferid\w*|v[ií]tima\w*|"
    r"atingiu|atingiram|destruiu|destrui\w*|arras\w*|incendi\w*|"
    r"balea\w*|tiroteio|disparos|abriu fogo|abrir fogo|"
    r"emboscad\w*|sequestr\w*|rapt\w*|decapit\w*|massacr\w*"
    r")\b"
)

# Contexto de ataque ocorrido, ainda que sem mortos/feridos explícitos
ATTACK_EVENT_CONTEXT_PAT = (
    r"\b("
    r"atentad\w*|ataque\w*|atac\w*|bombarde\w*|carro-bomba|homem-bomba|mulher-bomba|"
    r"explos[aã]o em|explos[aã]o no|explos[aã]o na|explos[aã]o junto|"
    r"ataque suicida|ataque armado|incurs[aã]o armada|emboscada|"
    r"granada|morteiro|foguet[aã]o|rocket|raide|raid"
    r")\b"
)

# Marcadores tipicamente não consumados / pre-evento que ajudam a evitar falsos positivos
PRE_EVENT_GUARD_PAT = (
    r"\b("
    r"suspeit\w*|alegad\w*|acusad\w*|investiga\w*|julgament\w*|tribunal|detido|deten[cç][aã]o|"
    r"pris[aã]o|condenad\w*|processo|inqu[eé]rito|"
    r"amea[cç]a\w*|plano\w*|planea\w*|planej\w*|tentativ\w*|falh\w*|frustrad\w*"
    r")\b"
)


def classify_violence_from_title(title: str):
    """
    Devolve:
      violence_detected_title (0/1)
      violence_type
      violence_outcome
      attack_detected_title (0/1)

    Lógica:
    - Falhado/Tentativa e Ameaça/Plano continuam prioritários.
    - "Consumado" passa a incluir:
        * marcadores fortes de evento ocorrido
        * qualquer tipo de violência reconhecido sem sinais de plano/falha
        * contexto claro de ataque/atentado sem sinais de plano/falha
    """
    t = normalize_text(title)

    if not t:
        return 0, "Outro/Não especificado", "Desconhecido", 0

    vtype = "Outro/Não especificado"
    type_matched = False
    for label, pat in VIOLENCE_TYPE_PATTERNS:
        if re.search(pat, t, flags=re.IGNORECASE):
            vtype = label
            type_matched = True
            break

    has_failed = bool(re.search(OUTCOME_FAILED_PAT, t, flags=re.IGNORECASE))
    has_attempt = bool(re.search(OUTCOME_ATTEMPT_PAT, t, flags=re.IGNORECASE))
    has_threat_plan = bool(re.search(OUTCOME_THREAT_PLAN_PAT, t, flags=re.IGNORECASE))
    has_occurred_strict = bool(re.search(OUTCOME_OCCURRED_STRICT_PAT, t, flags=re.IGNORECASE))
    has_attack_context = bool(re.search(ATTACK_EVENT_CONTEXT_PAT, t, flags=re.IGNORECASE))
    has_pre_event_guard = bool(re.search(PRE_EVENT_GUARD_PAT, t, flags=re.IGNORECASE))

    attack_detected_title = int(type_matched or has_occurred_strict or has_attack_context)

    if has_failed or has_attempt:
        outcome = "Falhado/Tentativa"
    elif has_threat_plan and not has_occurred_strict and not type_matched:
        outcome = "Ameaça/Plano"
    elif has_occurred_strict:
        outcome = "Consumado"
    elif type_matched and not has_threat_plan:
        outcome = "Consumado"
    elif has_attack_context and not has_pre_event_guard:
        outcome = "Consumado"
    elif has_threat_plan:
        outcome = "Ameaça/Plano"
    else:
        outcome = "Desconhecido"

    violence_detected_title = int(
        type_matched
        or has_occurred_strict
        or has_attack_context
        or has_failed
        or has_attempt
        or has_threat_plan
    )

    return violence_detected_title, vtype, outcome, attack_detected_title

# --------------------------------------------------
# DB helpers
# --------------------------------------------------
@st.cache_data(show_spinner=False)
def has_table(table_name: str) -> bool:
    return con.execute(
        "SELECT COUNT(*) FROM information_schema.tables WHERE table_name = ?",
        [table_name]
    ).fetchone()[0] > 0


@st.cache_data(show_spinner=False)
def date_bounds():
    return con.execute(
        "SELECT min(CAST(date_norm AS DATE)), max(CAST(date_norm AS DATE)) FROM news_analysis"
    ).fetchone()


@st.cache_data(show_spinner=False)
def distinct_fonte_lusa():
    return [r[0] for r in con.execute("""
        SELECT DISTINCT fonte
        FROM news_analysis
        WHERE fonte IS NOT NULL AND lower(fonte) LIKE '%lusa%'
        ORDER BY fonte
    """).fetchall()]


@st.cache_data(show_spinner=False)
def distinct_values_lusa(col: str):
    allowed = {"action_type", "place_country"}
    if col not in allowed:
        return []
    return [r[0] for r in con.execute(f"""
        SELECT DISTINCT {col}
        FROM news_analysis
        WHERE {col} IS NOT NULL
          AND lower(coalesce(fonte,'')) LIKE '%lusa%'
        ORDER BY {col}
    """).fetchall()]


@st.cache_data(show_spinner=False)
def gtd_columns_daily() -> set[str]:
    rows = con.execute("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_name = 'gtd_daily_country'
    """).fetchall()
    return {r[0] for r in rows}


@st.cache_data(show_spinner=False)
def load_lusa_base(
    date_from,
    date_to,
    fonte_sel_tuple,
    country_sel_tuple,
    action_type_sel_tuple,
    search_text: str,
    search_mode: str,
) -> pd.DataFrame:
    where = [
        "lower(coalesce(fonte,'')) LIKE '%lusa%'",
        "CAST(date_norm AS DATE) BETWEEN ? AND ?",
    ]
    params = [date_from, date_to]

    if fonte_sel_tuple:
        where.append("fonte IN (" + ",".join(["?"] * len(fonte_sel_tuple)) + ")")
        params.extend(list(fonte_sel_tuple))

    if country_sel_tuple:
        where.append("place_country IN (" + ",".join(["?"] * len(country_sel_tuple)) + ")")
        params.extend(list(country_sel_tuple))

    if action_type_sel_tuple:
        where.append("action_type IN (" + ",".join(["?"] * len(action_type_sel_tuple)) + ")")
        params.extend(list(action_type_sel_tuple))

    if search_text:
        target_expr = "lower(coalesce(titulo,'')) || ' ' || lower(coalesce(keyword_text_sources,''))"

        if search_mode == "Texto (contém)":
            where.append(f"{target_expr} LIKE ?")
            params.append(f"%{search_text.lower()}%")
        else:
            where.append(f"regexp_matches({target_expr}, ?)")
            params.append(search_text)

    q = f"""
        SELECT
          CAST(date_norm AS DATE) AS d,
          date_norm,
          titulo,
          fonte,
          place_country,
          place_country_iso,
          action_type,
          action_type_prev,
          keyword_text_sources
        FROM news_analysis
        WHERE {" AND ".join(where)}
    """
    df = con.execute(q, params).df()

    # --------------------------------------------------
    # DEDUPLICA??O ? uma not?cia por dia + t?tulo
    # --------------------------------------------------
    if not df.empty:
        df["_titulo_dedup"] = normalize_series(df["titulo"])
        df = df.drop_duplicates(
            subset=["d", "_titulo_dedup"],
            keep="first",
        ).copy()
        df = df.drop(columns=["_titulo_dedup"])

    return df

# --------------------------------------------------
# Construção lazy de padrões experimentais
# --------------------------------------------------
@st.cache_data(show_spinner=False)
def distinct_place_countries_cached() -> list[str]:
    return distinct_values_lusa("place_country")


def build_country_title_patterns(country_names: list[str]) -> list[tuple[str, re.Pattern]]:
    patterns = []
    seen = set()

    for country in country_names:
        c = str(country).strip()
        if not c:
            continue

        c_norm = normalize_text(c)
        if c_norm in seen:
            continue
        seen.add(c_norm)

        aliases = [c_norm]
        if c in COUNTRY_TITLE_ALIASES:
            aliases.extend(normalize_text(a) for a in COUNTRY_TITLE_ALIASES[c])

        aliases_clean = []
        seen_alias = set()
        for a in aliases:
            if a and a not in seen_alias:
                seen_alias.add(a)
                aliases_clean.append(a)

        pattern_txt = r"|".join(rf"\b{re.escape(a)}\b" for a in aliases_clean if a)
        if pattern_txt:
            patterns.append((c, re.compile(pattern_txt, re.IGNORECASE)))

    patterns.sort(key=lambda x: len(x[0]), reverse=True)
    return patterns


@st.cache_data(show_spinner=False)
def get_country_title_patterns():
    return build_country_title_patterns(distinct_place_countries_cached())


def detect_countries_in_title(title: str) -> list[str]:
    t = normalize_text(title)
    if not t:
        return []

    found = []
    for country, pat in get_country_title_patterns():
        if pat.search(t):
            found.append(country)
    return found


def build_detected_country_summary(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty or "titulo" not in df.columns:
        return pd.DataFrame(columns=["pais_detetado_no_titulo", "titulos", "percentagem"])

    rows = []
    for title in df["titulo"].fillna("").astype(str):
        countries = detect_countries_in_title(title)
        for c in countries:
            rows.append(c)

    if not rows:
        return pd.DataFrame(columns=["pais_detetado_no_titulo", "titulos", "percentagem"])

    out = (
        pd.DataFrame({"pais_detetado_no_titulo": rows})
        .groupby("pais_detetado_no_titulo", as_index=False)
        .size()
        .rename(columns={"size": "titulos"})
        .sort_values(["titulos", "pais_detetado_no_titulo"], ascending=[False, True])
    )

    total = int(out["titulos"].sum())
    out["percentagem"] = (out["titulos"] / total * 100).round(2)
    return out

# --------------------------------------------------
# UI — Header
# --------------------------------------------------
st.title("Lusa análise")
st.caption(
    "Principal: place_country. Experimental: deteção automática de países no título em tabela separada."
)

# --------------------------------------------------
# SIDEBAR — FILTROS BASE
# --------------------------------------------------
with st.sidebar:
    _, centro, _ = st.columns([1, 2, 1])
    with centro:
        st.image("logo_lxt.png", width=75)

    st.header("Filtros")

    dmin, dmax = date_bounds()
    dmin_dt = pd.to_datetime(dmin) if dmin is not None else pd.to_datetime("1987-01-01")
    dmax_dt = pd.to_datetime(dmax) if dmax is not None else pd.to_datetime("today")

    date_value = st.date_input("Intervalo temporal", value=(dmin_dt.date(), dmax_dt.date()))
    date_from, date_to = (date_value if isinstance(date_value, (tuple, list)) else (date_value, date_value))

    st.subheader("Origem editorial (Lusa)")
    fonte_sel = st.multiselect("Fonte (apenas Lusa)", distinct_fonte_lusa())

    st.subheader("Filtros textuais")

    st.markdown("**Pesquisa de palavras**")
    search_text = st.text_input("Palavra/frase", value="").strip()

    search_mode = st.radio(
        "Modo de pesquisa",
        ["Texto (contém)", "Regex"],
        horizontal=True,
        help="Use 'Texto (contém)' para pesquisa simples. Use 'Regex' apenas para padrões avançados.",
    )

    if search_mode == "Regex":
        st.caption(
            "Regex = pesquisa por padrões. Exemplos: "
            "`terror\\w*` apanha terror, terrorismo, terrorista; "
            "`ataque|ameaça` apanha ataque ou ameaça."
        )

    st.markdown("**Famílias lexicais**")
    family_scope = st.radio(
        "Onde aplicar famílias lexicais",
        ["Só no título", "No título + keyword_text_sources"],
        horizontal=False,
        help="Define se as famílias lexicais são procuradas só no título ou também em keyword_text_sources."
    )

    fam_terror = st.checkbox("terror*", value=False)
    fam_radic = st.checkbox("radic*", value=False)
    fam_extrem = st.checkbox("extrem*", value=False)
    fam_rebel = st.checkbox("rebel*", value=False)
    fam_mode = st.radio("Combinação das famílias", ["Qualquer (OR)", "Todas (AND)"], horizontal=True)

    st.subheader("Classificação e conteúdo")
    action_type_sel = st.multiselect("Action type (principal)", distinct_values_lusa("action_type"))

    st.subheader("Autor — tipologia macro (por título)")
    author_sel = st.multiselect("Selecionar", AUTHOR_OPTIONS)

    st.subheader("Política (por título)")
    politica_scope = st.radio(
        "Definição",
        ["Nacional", "Externa", "Nacional + Externa"],
        horizontal=False,
    )
    politica_mode = st.radio(
        "Filtro",
        ["Todos", "Só política", "Excluir política"],
        horizontal=False,
    )

    st.subheader("Violência (detalhe) — por título")
    violence_mode = st.radio("Filtro", ["Todos", "Só violência", "Excluir violência"], horizontal=True)
    vtype_sel, vout_sel = [], []
    if violence_mode != "Todos":
        vtype_sel = st.multiselect("Tipo de ataque", VIOLENCE_TYPE_OPTIONS, default=[])
        vout_sel = st.multiselect("Resultado", VIOLENCE_OUTCOME_OPTIONS, default=[])

    st.subheader("Tabela")
    max_rows = st.slider("Máximo de linhas (tabela)", 100, 5000, 500, 100)

    st.subheader("GTD")
    compare_gtd = st.checkbox("Incluir GTD no gráfico (série adicional)", value=True)

    st.subheader("Países")
    country_sel = st.multiselect("Filtrar por país (place_country)", distinct_values_lusa("place_country"))
    st.caption("Os blocos abaixo usam os resultados atuais já filtrados.")

# --------------------------------------------------
# LUSA — carregar base + heurísticas
# --------------------------------------------------
country_sel_combined = sorted(set(country_sel))

df_base = load_lusa_base(
    date_from,
    date_to,
    tuple(fonte_sel),
    tuple(country_sel_combined),
    tuple(action_type_sel),
    search_text,
    search_mode,
)

st.subheader("Resultados — Lusa (títulos)")

if df_base is None or df_base.empty:
    st.info("Sem resultados com estes filtros.")
    df_pre_countryfreq = pd.DataFrame()
else:
    df = df_base.copy()

    df = ensure_columns(
        df,
        {
            "titulo": "",
            "fonte": "",
            "place_country": "",
            "place_country_iso": "",
            "action_type": "",
            "action_type_prev": "",
            "keyword_text_sources": "",
            "date_norm": None,
            "d": None,
        }
    )

    if "titulo_norm" not in df.columns:
        df["titulo_norm"] = normalize_series(df["titulo"])

    df["author_macro"] = df["titulo_norm"].apply(macro_author_from_title)

    df["is_politica_nacional"] = df["titulo_norm"].str.contains(POLITICA_NAC_PAT, na=False).astype("int8")
    df["is_politica_externa"] = df["titulo_norm"].str.contains(POLITICA_EXT_PAT, na=False).astype("int8")
    df["is_politica"] = ((df["is_politica_nacional"] == 1) | (df["is_politica_externa"] == 1)).astype("int8")

    df = add_family_columns(df)

    v = df["titulo"].apply(classify_violence_from_title)
    df["violence_detected_title"] = [a for a, b, c, d in v]
    df["violence_type"] = [b for a, b, c, d in v]
    df["violence_outcome"] = [c for a, b, c, d in v]
    df["attack_detected"] = [d for a, b, c, d in v]

    if USE_CONSERVATIVE_VIOLENCE_FALLBACK:
        df["violence_detected"] = (
            (df["violence_detected_title"] == 1)
            | (df["attack_detected"] == 1)
            | (df["action_type_prev"].fillna("").astype(str).str.lower() == "violência")
            | (df["action_type_prev"].fillna("").astype(str).str.lower() == "violencia")
        ).astype(int)
    else:
        df["violence_detected"] = (
            (df["violence_detected_title"] == 1)
            | (df["attack_detected"] == 1)
        ).astype(int)

    if author_sel:
        df = df[df["author_macro"].isin(author_sel)]

    if politica_scope == "Nacional":
        pol_col = "is_politica_nacional"
    elif politica_scope == "Externa":
        pol_col = "is_politica_externa"
    else:
        pol_col = "is_politica"

    if politica_mode == "Só política":
        df = df[df[pol_col] == 1]
    elif politica_mode == "Excluir política":
        df = df[df[pol_col] == 0]

    if family_scope == "Só no título":
        family_map = {
            "terror": "has_terror_title",
            "radic": "has_radic_title",
            "extrem": "has_extrem_title",
            "rebel": "has_rebel_title",
        }
        family_label_scope = "title"
    else:
        family_map = {
            "terror": "has_terror_kw",
            "radic": "has_radic_kw",
            "extrem": "has_extrem_kw",
            "rebel": "has_rebel_kw",
        }
        family_label_scope = "kw"

    fam_cols = []
    if fam_terror:
        fam_cols.append(family_map["terror"])
    if fam_radic:
        fam_cols.append(family_map["radic"])
    if fam_extrem:
        fam_cols.append(family_map["extrem"])
    if fam_rebel:
        fam_cols.append(family_map["rebel"])

    if fam_cols:
        mask_df = df[fam_cols].eq(1)
        if fam_mode == "Todas (AND)":
            df = df[mask_df.all(axis=1)]
        else:
            df = df[mask_df.any(axis=1)]

    if violence_mode == "Só violência":
        df = df[df["violence_detected"] == 1]
    elif violence_mode == "Excluir violência":
        df = df[df["violence_detected"] == 0]

    if violence_mode != "Todos":
        if vtype_sel:
            df = df[df["violence_type"].isin(vtype_sel)]
        if vout_sel:
            df = df[df["violence_outcome"].isin(vout_sel)]

    df["data"] = parse_date_series(df["date_norm"]).dt.strftime("%d/%m/%Y")
    df["ano"] = parse_date_series(df["date_norm"]).dt.year
    df["familias_aplicadas"] = build_family_labels(df, scope=family_label_scope)

    if search_text and not df.empty:
        text_blob = (
            df["titulo"].fillna("").astype(str)
            + " "
            + df["keyword_text_sources"].fillna("").astype(str)
        )

        if search_mode == "Texto (contém)":
            needle = search_text.lower()
            hay = text_blob.fillna("").astype(str).str.lower()
            titles_with = hay.str.contains(re.escape(needle), regex=True, na=False).sum()
            total_occ = hay.str.count(re.escape(needle)).sum()
        else:
            try:
                titles_with = text_blob.str.contains(search_text, regex=True, na=False).sum()
                total_occ = text_blob.str.count(search_text, regex=True).sum()
            except re.error:
                titles_with = 0
                total_occ = 0
                st.warning("Regex inválida. Corrige o padrão ou usa 'Texto (contém)'.")

        st.markdown("### Estatísticas da pesquisa")
        st.write(f"Títulos com a palavra/frase: {int(titles_with):,}".replace(",", "."))
        st.write(f"Total de ocorrências: {int(total_occ):,}".replace(",", "."))

    df_pre_countryfreq = df.copy()

# --------------------------------------------------
# SIDEBAR DINÂMICA — Top 20 países atuais
# --------------------------------------------------
df_paises_sidebar = build_country_summary(df_pre_countryfreq)

with st.sidebar:
    st.subheader("Top 20 países atuais")

    if df_paises_sidebar.empty:
        st.caption("Sem países identificados nos resultados atuais.")
        country_freq_sel = []
    else:
        top20_sidebar = df_paises_sidebar.head(20).copy()
        linhas_top20 = [
            f"{row.place_country} — {int(row.noticias):,}".replace(",", ".")
            for _, row in top20_sidebar.iterrows()
        ]
        st.caption("\n".join(linhas_top20))

        paises_ordenados_freq = df_paises_sidebar["place_country"].tolist()

        country_freq_sel = st.multiselect(
            "Selecionar países por frequência",
            options=paises_ordenados_freq,
            default=[],
            help="Lista ordenada dos países mais frequentes nos resultados atuais."
        )

# --------------------------------------------------
# Aplicar filtro extra de países por frequência
# --------------------------------------------------
if df_pre_countryfreq is None or df_pre_countryfreq.empty:
    df_final = pd.DataFrame()
else:
    df_final = df_pre_countryfreq.copy()
    if country_freq_sel:
        df_final = df_final[df_final["place_country"].isin(country_freq_sel)].copy()

# --------------------------------------------------
# TABELA PRINCIPAL
# --------------------------------------------------
if df_final is not None and not df_final.empty:
    df_show = df_final.sort_values("d", ascending=False).head(max_rows).copy()

    cols_show = [
        "data", "titulo", "fonte",
        "place_country", "place_country_iso",
        "is_politica", "is_politica_nacional", "is_politica_externa",
        "author_macro",
        "action_type",
        "familias_aplicadas",
        "attack_detected",
        "violence_detected", "violence_type", "violence_outcome",
    ]
    cols_show = [c for c in cols_show if c in df_show.columns]

    st.dataframe(df_show[cols_show], use_container_width=True)

# --------------------------------------------------
# RESUMO + EXPORT — Lusa filtrada
# --------------------------------------------------
total_lusa = 0 if df_final is None else len(df_final)

st.markdown(
    f"""
    <div style="
        margin-top: 8px;
        margin-bottom: 14px;
        padding: 12px 16px;
        border: 1px solid rgba(255,255,255,0.10);
        border-radius: 12px;
        background: rgba(255,255,255,0.03);
        width: fit-content;
        min-width: 220px;
    ">
        <div style="
            font-size: 0.82rem;
            color: rgba(255,255,255,0.70);
            margin-bottom: 2px;
        ">
            Total de notícias
        </div>
        <div style="
            font-size: 1.55rem;
            font-weight: 700;
            line-height: 1.1;
        ">
            {f"{total_lusa:,}".replace(",", ".")}
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

with st.expander("Exportar resultados — Lusa filtrada", expanded=False):
    if df_final is None or df_final.empty:
        st.info("Sem linhas para exportar.")
    else:
        df_export = df_final.sort_values(["ano", "titulo"], ascending=[True, True]).copy()

        st.download_button(
            "Descarregar CSV (Lusa completa)",
            data=df_to_csv_bytes(df_export),
            file_name=f"lusa_filtrado_{date_from}_{date_to}.csv",
            mime="text/csv",
        )

        df_export_excel = df_export[["ano", "titulo", "place_country"]].copy()
        df_export_excel["ano"] = pd.to_numeric(df_export_excel["ano"], errors="coerce")
        df_export_excel = df_export_excel.dropna(subset=["ano"])
        df_export_excel["ano"] = df_export_excel["ano"].astype(int)
        df_export_excel = df_export_excel.sort_values(["ano", "place_country", "titulo"])

        st.download_button(
            "Descarregar Excel (ano + título + país)",
            data=df_to_excel_bytes(df_export_excel, sheet_name="resultados"),
            file_name=f"lusa_resultados_legiveis_{date_from}_{date_to}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

# --------------------------------------------------
# PAÍSES MAIS REFERIDOS — Lusa (principal)
# --------------------------------------------------
st.subheader("Países mais referidos — Lusa")
st.caption("Base principal: coluna estruturada place_country.")

if df_final is None or df_final.empty:
    st.info("Sem resultados para calcular países mais referidos.")
else:
    top_n_paises = st.slider("Top países", 5, 50, 15, 5, key="top_paises_slider")

    df_paises = build_country_summary(df_final)

    if df_paises.empty:
        st.info("Não há países identificados nos resultados filtrados.")
    else:
        total_noticias_com_pais = int(df_paises["noticias"].sum())

        st.markdown(
            f"""
            <div style="
                margin-top: 8px;
                margin-bottom: 14px;
                padding: 12px 16px;
                border: 1px solid rgba(255,255,255,0.10);
                border-radius: 12px;
                background: rgba(255,255,255,0.03);
                width: fit-content;
                min-width: 260px;
            ">
                <div style="
                    font-size: 0.82rem;
                    color: rgba(255,255,255,0.70);
                    margin-bottom: 2px;
                ">
                    Notícias com país identificado
                </div>
                <div style="
                    font-size: 1.55rem;
                    font-weight: 700;
                    line-height: 1.1;
                ">
                    {f"{total_noticias_com_pais:,}".replace(",", ".")}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.dataframe(df_paises.head(top_n_paises), use_container_width=True)

        st.download_button(
            "Descarregar CSV — países mais referidos",
            data=df_to_csv_bytes(df_paises),
            file_name=f"paises_mais_referidos_{date_from}_{date_to}.csv",
            mime="text/csv",
        )

        st.download_button(
            "Descarregar Excel — países mais referidos",
            data=build_country_excel_bytes(df_paises),
            file_name=f"paises_mais_referidos_{date_from}_{date_to}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

# --------------------------------------------------
# PAÍSES DETETADOS AUTOMATICAMENTE NO TÍTULO — experimental
# --------------------------------------------------
st.subheader("Países detetados automaticamente no título — experimental")
st.caption(
    "Secção auxiliar. Não substitui place_country. Baseia-se em correspondência automática de nomes/aliases no título."
)

if df_final is None or df_final.empty:
    st.info("Sem resultados para deteção experimental de países no título.")
else:
    df_detected_countries = build_detected_country_summary(df_final)

    if df_detected_countries.empty:
        st.info("Não foram detetados países nos títulos com as regras atuais.")
    else:
        total_detected_titles = int(df_detected_countries["titulos"].sum())

        st.markdown(
            f"""
            <div style="
                margin-top: 8px;
                margin-bottom: 14px;
                padding: 12px 16px;
                border: 1px solid rgba(255,255,255,0.10);
                border-radius: 12px;
                background: rgba(255,255,255,0.03);
                width: fit-content;
                min-width: 300px;
            ">
                <div style="
                    font-size: 0.82rem;
                    color: rgba(255,255,255,0.70);
                    margin-bottom: 2px;
                ">
                    Deteções automáticas em títulos
                </div>
                <div style="
                    font-size: 1.55rem;
                    font-weight: 700;
                    line-height: 1.1;
                ">
                    {f"{total_detected_titles:,}".replace(",", ".")}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.dataframe(df_detected_countries, use_container_width=True)

        st.download_button(
            "Descarregar CSV — países detetados no título",
            data=df_to_csv_bytes(df_detected_countries),
            file_name=f"paises_detetados_titulo_{date_from}_{date_to}.csv",
            mime="text/csv",
        )

        st.download_button(
            "Descarregar Excel — países detetados no título",
            data=build_detected_country_excel_bytes(df_detected_countries),
            file_name=f"paises_detetados_titulo_{date_from}_{date_to}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

# --------------------------------------------------
# COMPARAÇÃO Lusa × GTD — dia–ISO
# --------------------------------------------------
st.subheader("Comparação Lusa × GTD — dia–país (ISO)")

st.caption(
    "Leitura da comparacao: a serie Lusa conta as noticias "
    "que correspondem aos filtros selecionados e cujo pais esta identificado. "
    "A serie GTD conta os atentados registados para o mesmo dia e pais. "
    "A coincidencia de data e pais nao confirma que uma noticia corresponda "
    "a um atentado especifico. Os resultados da Lusa variam consoante os filtros aplicados."
)

comp = pd.DataFrame()
comp_show = pd.DataFrame()
comp_cols = []

if df_final is None or df_final.empty:
    st.info("Sem linhas Lusa filtradas para comparar.")
elif not has_table("gtd_daily_country"):
    st.info("Tabela GTD (gtd_daily_country) não encontrada na base.")
else:
    tmp = df_final.copy()
    tmp = tmp[tmp["place_country_iso"].fillna("").astype(str).str.strip() != ""]

    lusa_agg = (
        tmp.groupby(["d", "place_country_iso"], as_index=False)
        .size()
        .rename(columns={"size": "lusa_n", "place_country_iso": "country_iso", "d": "dia"})
    )
    lusa_agg["dia"] = pd.to_datetime(lusa_agg["dia"]).dt.date

    iso_to_name = (
        tmp.assign(_iso=tmp["place_country_iso"].astype(str).str.strip())
        .loc[lambda x: x["_iso"] != "", ["_iso", "place_country"]]
        .groupby("_iso")["place_country"]
        .agg(lambda s: s.value_counts().index[0] if len(s) else "")
        .to_dict()
    )
    lusa_agg["country_name"] = lusa_agg["country_iso"].map(iso_to_name).fillna("")

    gtd_cols = gtd_columns_daily()
    has_country_txt = "country_txt" in gtd_cols

    if has_country_txt:
        gtd = con.execute("""
            SELECT
              event_date AS dia,
              country_iso,
              ANY_VALUE(country_txt) AS country_txt,
              SUM(gtd_events) AS gtd_events,
              SUM(gtd_kill) AS gtd_kill,
              SUM(gtd_wound) AS gtd_wound,
              SUM(gtd_failed) AS gtd_failed
            FROM gtd_daily_country
            WHERE event_date BETWEEN ? AND ?
            GROUP BY 1,2
        """, [date_from, date_to]).df()
    else:
        gtd = con.execute("""
            SELECT
              event_date AS dia,
              country_iso,
              SUM(gtd_events) AS gtd_events,
              SUM(gtd_kill) AS gtd_kill,
              SUM(gtd_wound) AS gtd_wound,
              SUM(gtd_failed) AS gtd_failed
            FROM gtd_daily_country
            WHERE event_date BETWEEN ? AND ?
            GROUP BY 1,2
        """, [date_from, date_to]).df()

    gtd["dia"] = pd.to_datetime(gtd["dia"]).dt.date

    comp = pd.merge(lusa_agg, gtd, how="outer", on=["dia", "country_iso"])
    comp["lusa_n"] = comp["lusa_n"].fillna(0).astype(int)
    for c in ["gtd_events", "gtd_kill", "gtd_wound", "gtd_failed"]:
        if c in comp.columns:
            comp[c] = safe_int_series(comp[c])

    comp["dif_lusa_minus_gtd"] = comp["lusa_n"] - comp.get("gtd_events", 0)
    comp = comp.sort_values(["dia", "country_iso"], ascending=[False, True])

    comp_show = comp.copy()
    comp_show["dia"] = pd.to_datetime(comp_show["dia"]).dt.strftime("%d/%m/%Y")

    comp_cols = ["dia", "country_iso", "country_name"]
    if has_country_txt:
        comp_cols.append("country_txt")
    comp_cols += ["lusa_n", "gtd_events", "gtd_kill", "gtd_wound", "gtd_failed", "dif_lusa_minus_gtd"]
    comp_cols = [c for c in comp_cols if c in comp_show.columns]

    st.dataframe(comp_show[comp_cols], use_container_width=True)

# --------------------------------------------------
# LEGENDA + EXPORT — Comparação
# --------------------------------------------------
if comp_show is not None and not comp_show.empty:
    st.caption(f"Total de linhas: {len(comp_show):,}".replace(",", "."))

    with st.expander("Exportar comparação Lusa × GTD", expanded=False):
        st.download_button(
            "Descarregar CSV (Comparação Lusa × GTD)",
            data=df_to_csv_bytes(comp_show[comp_cols]),
            file_name=f"comparacao_lusa_gtd_{date_from}_{date_to}.csv",
            mime="text/csv",
        )

# --------------------------------------------------
# GTD — LISTA EVENTO-A-EVENTO (gtd_raw)
# --------------------------------------------------
st.subheader("GTD — Lista de eventos (event-level)")

if not has_table("gtd_raw"):
    st.info("Tabela event-level não encontrada (gtd_raw).")
else:
    with st.expander("Ver e exportar GTD (event-level)", expanded=False):
        st.caption("Preview (Top N) na app. Para lista completa use export CSV (intervalo atual).")
        gtd_limit = st.slider("Preview (Top N)", 100, 20000, 2000, 100)

        has_dim_iso = has_table("dim_country_txt_to_iso")

        if has_dim_iso:
            q_preview = """
            WITH base AS (
              SELECT
                r.*,
                CASE
                  WHEN try_cast(r.iyear AS INTEGER) IS NOT NULL
                   AND try_cast(r.imonth AS INTEGER) BETWEEN 1 AND 12
                   AND try_cast(r.iday AS INTEGER) BETWEEN 1 AND 31
                  THEN make_date(
                    try_cast(r.iyear AS INTEGER),
                    try_cast(r.imonth AS INTEGER),
                    try_cast(r.iday AS INTEGER)
                  )
                  ELSE NULL
                END AS event_date
              FROM gtd_raw r
            )
            SELECT
              b.eventid,
              b.event_date,
              b.country_txt,
              m.country_iso,
              b.region_txt,
              b.city,
              b.attacktype1_txt,
              b.targtype1_txt,
              b.weaptype1_txt,
              b.success,
              b.suicide,
              b.nkill,
              b.nwound,
              b.summary
            FROM base b
            LEFT JOIN dim_country_txt_to_iso m
              ON lower(trim(b.country_txt)) = lower(trim(m.country_txt))
            WHERE b.event_date IS NOT NULL
              AND b.event_date BETWEEN ? AND ?
            ORDER BY b.event_date DESC, b.eventid
            LIMIT ?
            """
        else:
            q_preview = """
            WITH base AS (
              SELECT
                r.*,
                CASE
                  WHEN try_cast(r.iyear AS INTEGER) IS NOT NULL
                   AND try_cast(r.imonth AS INTEGER) BETWEEN 1 AND 12
                   AND try_cast(r.iday AS INTEGER) BETWEEN 1 AND 31
                  THEN make_date(
                    try_cast(r.iyear AS INTEGER),
                    try_cast(r.imonth AS INTEGER),
                    try_cast(r.iday AS INTEGER)
                  )
                  ELSE NULL
                END AS event_date
              FROM gtd_raw r
            )
            SELECT
              b.eventid,
              b.event_date,
              b.country_txt,
              b.region_txt,
              b.city,
              b.attacktype1_txt,
              b.targtype1_txt,
              b.weaptype1_txt,
              b.success,
              b.suicide,
              b.nkill,
              b.nwound,
              b.summary
            FROM base b
            WHERE b.event_date IS NOT NULL
              AND b.event_date BETWEEN ? AND ?
            ORDER BY b.event_date DESC, b.eventid
            LIMIT ?
            """

        df_gtd_prev = con.execute(q_preview, [date_from, date_to, gtd_limit]).df()
        st.dataframe(df_gtd_prev, use_container_width=True)

        if st.button("Exportar GTD event-level (CSV) — intervalo atual"):
            if has_dim_iso:
                q_export = """
                WITH base AS (
                  SELECT
                    r.*,
                    CASE
                      WHEN try_cast(r.iyear AS INTEGER) IS NOT NULL
                       AND try_cast(r.imonth AS INTEGER) BETWEEN 1 AND 12
                       AND try_cast(r.iday AS INTEGER) BETWEEN 1 AND 31
                      THEN make_date(
                        try_cast(r.iyear AS INTEGER),
                        try_cast(r.imonth AS INTEGER),
                        try_cast(r.iday AS INTEGER)
                      )
                      ELSE NULL
                    END AS event_date
                  FROM gtd_raw r
                )
                SELECT
                  b.*,
                  m.country_iso
                FROM base b
                LEFT JOIN dim_country_txt_to_iso m
                  ON lower(trim(b.country_txt)) = lower(trim(m.country_txt))
                WHERE b.event_date IS NOT NULL
                  AND b.event_date BETWEEN ? AND ?
                ORDER BY b.event_date DESC, b.eventid
                """
            else:
                q_export = """
                WITH base AS (
                  SELECT
                    r.*,
                    CASE
                      WHEN try_cast(r.iyear AS INTEGER) IS NOT NULL
                       AND try_cast(r.imonth AS INTEGER) BETWEEN 1 AND 12
                       AND try_cast(r.iday AS INTEGER) BETWEEN 1 AND 31
                      THEN make_date(
                        try_cast(r.iyear AS INTEGER),
                        try_cast(r.imonth AS INTEGER),
                        try_cast(r.iday AS INTEGER)
                      )
                      ELSE NULL
                    END AS event_date
                  FROM gtd_raw r
                )
                SELECT
                  b.*
                FROM base b
                WHERE b.event_date IS NOT NULL
                  AND b.event_date BETWEEN ? AND ?
                ORDER BY b.event_date DESC, b.eventid
                """

            df_gtd_export = con.execute(q_export, [date_from, date_to]).df()
            st.download_button(
                "Descarregar CSV GTD (event-level)",
                data=df_to_csv_bytes(df_gtd_export),
                file_name=f"gtd_event_level_{date_from}_{date_to}.csv",
                mime="text/csv",
            )
            st.write(f"Linhas exportadas: {len(df_gtd_export):,}".replace(",", "."))

# --------------------------------------------------
# GRÁFICO — Evolução temporal (Lusa + GTD)
# --------------------------------------------------
st.subheader("Gráfico — Evolução temporal (Lusa + GTD)")

gran = st.radio("Agregação", ["Dia", "Mês", "Ano"], horizontal=True, key="gran_radio")


def bucket_dates(dt: pd.Series) -> pd.Series:
    dt = pd.to_datetime(dt, errors="coerce")
    if gran == "Mês":
        return dt.dt.to_period("M").dt.to_timestamp()
    if gran == "Ano":
        return dt.dt.to_period("Y").dt.to_timestamp()
    return dt.dt.floor("D")


df_chart = pd.DataFrame(columns=["t", "n", "serie"])

if df_final is not None and not df_final.empty:
    tmp = df_final.copy()
    tmp["t"] = bucket_dates(tmp["d"])
    ts_lusa = tmp.groupby("t", as_index=False).size().rename(columns={"size": "n"})
    ts_lusa["serie"] = "Lusa"
    df_chart = pd.concat([df_chart, ts_lusa], ignore_index=True)

if compare_gtd and has_table("gtd_daily_country"):
    ts_gtd = con.execute("""
        SELECT event_date AS t, SUM(gtd_events) AS n
        FROM gtd_daily_country
        WHERE event_date BETWEEN ? AND ?
        GROUP BY 1
        ORDER BY 1
    """, [date_from, date_to]).df()

    if not ts_gtd.empty:
        ts_gtd["t"] = bucket_dates(ts_gtd["t"])
        ts_gtd = ts_gtd.groupby("t", as_index=False)["n"].sum()
        ts_gtd["serie"] = "GTD"
        df_chart = pd.concat([df_chart, ts_gtd], ignore_index=True)

if df_chart.empty:
    st.info("Sem pontos para desenhar no gráfico com os filtros atuais.")
else:
    df_chart["t"] = pd.to_datetime(df_chart["t"], errors="coerce")
    df_chart["n"] = pd.to_numeric(df_chart["n"], errors="coerce").fillna(0).astype(int)

    chart = (
        alt.Chart(df_chart)
        .mark_line(point=alt.OverlayMarkDef(filled=True, size=40))
        .encode(
            x=alt.X("t:T", title="Tempo"),
            y=alt.Y("n:Q", title="Contagem"),
            color=alt.Color("serie:N", title="Série"),
            tooltip=[alt.Tooltip("serie:N"), alt.Tooltip("t:T"), alt.Tooltip("n:Q")],
        )
        .interactive()
    )
    st.altair_chart(chart, use_container_width=True)

    st.markdown("#### Exportar gráfico para Excel")

    df_chart_export = df_chart.copy()
    df_chart_export["t"] = pd.to_datetime(df_chart_export["t"], errors="coerce")
    df_chart_export["ano"] = df_chart_export["t"].dt.year

    df_chart_year = (
        df_chart_export.groupby(["ano", "serie"], as_index=False)["n"]
        .sum()
        .pivot(index="ano", columns="serie", values="n")
        .fillna(0)
        .reset_index()
    )

    for col in ["Lusa", "GTD"]:
        if col not in df_chart_year.columns:
            df_chart_year[col] = 0

    df_chart_year = df_chart_year[["ano", "Lusa", "GTD"]].copy()
    df_chart_year["ano"] = pd.to_numeric(df_chart_year["ano"], errors="coerce")
    df_chart_year = df_chart_year.dropna(subset=["ano"])
    df_chart_year["ano"] = df_chart_year["ano"].astype(int)
    df_chart_year["Lusa"] = pd.to_numeric(df_chart_year["Lusa"], errors="coerce").fillna(0).astype(int)
    df_chart_year["GTD"] = pd.to_numeric(df_chart_year["GTD"], errors="coerce").fillna(0).astype(int)
    df_chart_year = df_chart_year.sort_values("ano")

    st.download_button(
        "Exportar dados + gráfico para Excel",
        data=build_chart_excel_bytes(df_chart_year, title="Evolução temporal — Lusa vs GTD"),
        file_name=f"grafico_lusa_gtd_ano_{date_from}_{date_to}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

